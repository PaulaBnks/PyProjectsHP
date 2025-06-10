from simple_salesforce import SalesforceLogin, Salesforce
from google.cloud import bigquery
from google.cloud.bigquery import SchemaField
import json
import google.generativeai as genai
import requests
import base64
from docx import Document
from io import BytesIO
from datetime import datetime
from google.oauth2 import service_account
from bs4 import BeautifulSoup
import re
import os
import dateparser  # For smart date parsing

# --- CONFIGURATION ---
run_on_git_actions = 'No'

# --- AUTHENTICATION ---
if run_on_git_actions == 'Yes':
    sf = Salesforce(
        username=os.environ["username"],
        password=os.environ["pwd"],
        security_token=os.environ["token"],
        client_id='python'
    )
    session_id, instance = SalesforceLogin(
        username=os.environ["username"],
        password=os.environ["pwd"],
        security_token=os.environ["token"]
    )
else:
    print("🔐 Logging into Salesforce...")
    login_info = json.load(open(r"E:\Software\loginpd.json"))
    username = login_info['username']
    password = login_info['password']
    security_token = login_info['security_token']
    session_id, instance = SalesforceLogin(username=username, password=password, security_token=security_token)
    sf = Salesforce(instance=instance, session_id=session_id)
    print("✅ Salesforce login successful.")

# Authenticate Gemini
print("🔐 Loading Gemini API key...")
with open(r"E:\Software\gemini_prod_key.txt", "r") as file:
    api_key = file.read().strip()
genai.configure(api_key=api_key)
embedding_model = genai.GenerativeModel("models/embedding-001")
print("✅ Gemini model initialized.")

# Authenticate BigQuery
credentials = service_account.Credentials.from_service_account_file(
    r"E:\Software\python-project-403413-480829366e0a.json"
)
bq_client = bigquery.Client(project='southern-coda-233109', credentials=credentials)
print("✅ BigQuery login successful.")


# --- HELPER FUNCTIONS ---

def extract_docx_text_from_bytes(file_bytes):
    """Extract text from Word document bytes."""
    try:
        doc = Document(BytesIO(file_bytes))
        return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    except Exception as e:
        print(f"⚠️ DOCX could not be processed: {e}")
        return None


BASE64_CHARS = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789+/"

def extract_note_content(version, session_id, instance):
    version_id = version['Id']
    filetype = version.get('FileType')
    url = f"https://{instance}/services/data/v62.0/sobjects/ContentVersion/{version_id}?fields=CreatedDate"
    headers = {'Authorization': f'Bearer {session_id}'}
    meta_response = requests.get(url, headers=headers)
    created_date = meta_response.json().get('CreatedDate', '') if meta_response.status_code == 200 else ''

    download_url = f"https://{instance}/services/data/v62.0/sobjects/ContentVersion/{version_id}/VersionData"
    content_response = requests.get(download_url, headers=headers)

    if content_response.status_code != 200:
        return None, created_date

    content_bytes = content_response.content
    if filetype in ['PLAINTEXT', 'HTML']:
        content = content_bytes.decode('utf-8', errors='ignore')
        return content.strip(), created_date
    elif filetype == 'SNOTE':
        decoded = safe_base64_decode(content_bytes.decode('utf-8', errors='ignore'))
        return decoded.strip(), created_date
    elif filetype == 'WORD_X':
        return extract_docx_text_from_bytes(content_bytes), created_date
    return None, created_date

def safe_base64_decode(data, account_id=None):
    """Attempts to decode base64 content safely."""
    if not data or (isinstance(data, str) and data.strip() == ''):
        return ''
    if isinstance(data, bytes):
        try:
            data = data.decode('utf-8')
        except UnicodeDecodeError:
            try:
                data = data.decode('latin1')
            except Exception as e:
                print("      ⚠️ Unable to decode byte stream.")
                return "[Unreadable note content: Invalid byte stream]"
    data = data.strip()
    if len(data) >= 2 and all(c in BASE64_CHARS for c in data[:50]):
        pass
    else:
        return data
    missing_padding = len(data) % 4
    if missing_padding:
        data += '=' * (4 - missing_padding)
    try:
        decoded_bytes = base64.b64decode(data, validate=False, altchars=None)
        decoded_text = decoded_bytes.decode('utf-8', errors='replace')
        if decoded_text.count('\uFFFD') > len(decoded_text) // 4:
            raise ValueError("Too many invalid characters after decode")
        return decoded_text
    except Exception as e:
        return f"[Unreadable note content: {data[:100]}...]"


def strip_html(text):
    """Remove HTML tags and decode HTML entities."""
    if not text:
        return ""
    try:
        soup = BeautifulSoup(text, "html.parser")
        clean_text = soup.get_text(separator="\n", strip=True)
        return clean_text
    except Exception as e:
        print(f"⚠️ Error stripping HTML: {e}")
        return text.strip()


def preprocess_note_content(text):
    """
    Removes unhelpful lines like Avoma URLs and JavaScript placeholders.
    Keeps actual meeting content.
    """
    lines = text.splitlines()
    cleaned_lines = []

    exclude_patterns = [
        r"^Avoma Meeting\s*$",
        r"https?://",
        r"^You need to enable JavaScript to run this app.$"
    ]

    for line in lines:
        line_clean = line.strip()
        if not line_clean:
            continue  # Skip empty lines

        excluded = any(re.search(pattern, line_clean) for pattern in exclude_patterns)
        if not excluded:
            cleaned_lines.append(line_clean)

    return "\n".join(cleaned_lines)


def normalize_dates(text, base_date=None):
    """
    Converts informal or relative dates to [DD.MM.YYYY, HH:MM].
    Uses base_date (meeting start time) as reference.
    """

    def replace_match(match):
        phrase = match.group(0)
        parsed = dateparser.parse(phrase, settings={'RELATIVE_BASE': base_date})
        if parsed:
            return f"[{parsed.strftime('%d.%m.%Y, %H:%M')}]"
        return phrase

    pattern = r"(next \w+ at \d{1,2}:\d{2})|(\w+day at \d{1,2}:\d{2})|(tomorrow at \d{1,2}:\d{2})|(\d{1,2}:\d{2}\s*[APMapm]+\s*\w+\s*\d{1,2},\s*\d{4})"
    return re.sub(pattern, replace_match, text, flags=re.IGNORECASE)


def convert_salesforce_time(dt_str):
    """Convert Salesforce datetime string to BigQuery-compatible format."""
    try:
        dt = datetime.strptime(dt_str, "%Y-%m-%dT%H:%M:%S.%f%z")
    except ValueError:
        try:
            dt = datetime.strptime(dt_str, "%Y-%m-%dT%H:%M:%S%z")
        except ValueError:
            try:
                dt = datetime.strptime(dt_str.split('.')[0], "%Y-%m-%dT%H:%M:%S")
            except Exception as e:
                print(f"⚠️ Could not parse date: {dt_str}")
                return None
    return dt


def get_embedding(text):
    result = genai.embed_content(
        model="models/embedding-001",
        content=text,
        task_type="retrieval_document"
    )
    return result["embedding"]


def save_note_embedding(account_id, version_id, raw_text, clean_text, created_date, embedding):
    table_id = "southern-coda-233109.import.salesforce_embeddings"

    # Convert timestamp to BQ-friendly format
    bq_created_date = created_date.strftime("%Y-%m-%d %H:%M:%S") if created_date else None

    schema = [
        SchemaField("id", "STRING"),
        SchemaField("account_id", "STRING"),
        SchemaField("object_type", "STRING"),
        SchemaField("created_date", "TIMESTAMP"),
        SchemaField("raw_text", "STRING"),
        SchemaField("clean_text", "STRING"),
        SchemaField("embedding", "FLOAT64", mode="REPEATED")
    ]

    rows_to_insert = [{
        "id": version_id,
        "account_id": account_id,
        "object_type": "Note",
        "created_date": bq_created_date,
        "raw_text": raw_text,
        "clean_text": clean_text,
        "embedding": embedding
    }]

    errors = bq_client.insert_rows(table_id, rows_to_insert, selected_fields=schema)

    if errors:
        print(f"❌ Error inserting row: {errors}")
    else:
        print(f"✅ Stored embedding for note {version_id}")


# --- MAIN SCRIPT LOGIC ---
accountid = '00109000013HoXkAAK'
print(f"\n📡 Fetching notes for Account: {accountid}")

note_embeddings = []

try:
    linked_docs = sf.query(f"SELECT ContentDocumentId FROM ContentDocumentLink WHERE LinkedEntityId = '{accountid}'")['records']
    doc_ids = [d['ContentDocumentId'] for d in linked_docs]

    for doc_id in doc_ids:
        versions = sf.query(f"SELECT Id, FileType, ContentDocument.CreatedBy.Name, ContentDocument.CreatedDate FROM ContentVersion WHERE ContentDocumentId = '{doc_id}' ORDER BY CreatedDate DESC LIMIT 1")['records']
        for version in versions:
            content, created_date_str = extract_note_content(version, session_id, instance)
            if content:
                created_date = convert_salesforce_time(created_date_str)
                version_id = version['Id']
                note_embeddings.append((content, created_date, version_id))

except Exception as e:
    print(f"⚠️ Error while fetching notes: {e}")


print(f"\n🧠 Generating embeddings for {len(note_embeddings)} notes...\n")

for idx, (text, created_date, version_id) in enumerate(note_embeddings):
    try:
        print(f"📌 Note {idx + 1}: {text[:100]}... (ID: {version_id})")

        # Step 0: Strip HTML if present
        html_stripped = strip_html(text)

        # Step 1: Remove boilerplate lines
        clean_text = preprocess_note_content(html_stripped)

        # Step 2: Normalize dates using meeting date as reference
        enriched_text = normalize_dates(clean_text, base_date=created_date)

        # Step 3: Generate embedding only if content is not empty
        if enriched_text.strip():
            try:
                embedding = get_embedding(enriched_text)
                print(f"🔢 Embedding (first 10 values): {embedding[:10]}\n")
            except Exception as e:
                print(f"❌ Failed to generate embedding: {e}")
                embedding = None
        else:
            print("🟡 Skipping embedding: Cleaned text is empty after preprocessing.")
            with open("empty_notes_debug.log", "a", encoding="utf-8") as f:
                f.write(f"\n--- EMPTY NOTE ID: {version_id} ---\n")
                f.write(f"Raw Text:\n{text}\n\nCleaned Text:\n{enriched_text}\n")

        # Step 4: Save to BigQuery
        save_note_embedding(accountid, version_id, text, enriched_text, created_date, embedding)

    except Exception as e:
        print(f"❌ Failed to process note {idx+1}: {e}")